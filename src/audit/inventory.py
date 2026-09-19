"""Generate offline source inventories without fitting scientific models."""

import json
from pathlib import Path
from xml.etree import ElementTree

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.acquisition.download import ROOT, MANIFEST, sha256


def require_unique(frame: pd.DataFrame, keys: list[str]) -> None:
    """Require nonmissing identifiers unique at the declared observation grain."""
    if frame[keys].isna().any().any() or frame.duplicated(keys).any():
        raise ValueError(f'Missing or duplicate keys: {keys}')


def align_samples(metadata: pd.DataFrame, sample_ids: list[str]) -> pd.DataFrame:
    """Align by explicit sample ID; reject missing, extra or duplicate inputs."""
    require_unique(metadata, ['ODCF_name'])
    if len(sample_ids) != len(set(sample_ids)):
        raise ValueError('Duplicate count sample IDs')
    if set(sample_ids) != set(metadata.ODCF_name):
        raise ValueError('Count/metadata sample sets differ')
    return metadata.set_index('ODCF_name').loc[sample_ids].reset_index()


def verify_manifest(prefix: str = "") -> int:
    records = [r for r in json.loads(MANIFEST.read_text()) if r['local_filename'].startswith(prefix)]
    if not records:
        raise ValueError('No provenance records match requested scope')
    for record in records:
        path = ROOT / record['local_filename']
        if not path.exists() or sha256(path) != record['sha256']:
            raise ValueError(f"Input integrity failure: {record['local_filename']}")
    return len(records)


def audit_schmidt() -> dict:
    base = ROOT / 'data/raw/schmidt/github/raw_data/IlseID29538'
    meta = pd.read_excel(base / 'IlseID29538_metadata.xlsx')
    paths = sorted((base / 'featureCounts').glob('*.tsv'))
    aligned = align_samples(meta, [p.name.split('_OE')[0] for p in paths])
    aligned.to_csv(ROOT / 'metadata/schmidt_samples.csv', index=False)
    dictionary = [{'original_column': column, 'dtype': str(meta[column].dtype),
                   'missing_n': int(meta[column].isna().sum()),
                   'unique_nonmissing_n': int(meta[column].nunique()),
                   'definition_status': 'source header; semantic review required'}
                  for column in meta.columns]
    pd.DataFrame(dictionary).to_csv(ROOT / 'metadata/schmidt_data_dictionary.csv', index=False)
    groups = aligned.groupby(['Cell_line', 'Batch', 'Condition_simple']).size().reset_index(name='libraries_n')
    groups.to_csv(ROOT / 'metadata/schmidt_design_counts.csv', index=False)
    reference = None
    counts = []
    for path in paths:
        table = pd.read_csv(path, sep='\t', usecols=['gene_id', 'num_reads', 'name'])
        require_unique(table, ['gene_id'])
        values = table.num_reads.to_numpy()
        if not (np.isfinite(values).all() and (values >= 0).all() and (values == np.floor(values)).all()):
            raise ValueError(f'Invalid count scale: {path.name}')
        genes = table.gene_id.tolist()
        if reference is None:
            reference = genes
        elif genes != reference:
            raise ValueError(f'Gene order/universe differs: {path.name}')
        counts.append({'sample_id': path.name.split('_OE')[0], 'genes_n': len(table),
                       'library_count_sum': int(values.sum()), 'zero_genes_n': int((values == 0).sum()),
                       'count_column': 'num_reads', 'integer_nonnegative': True})
    pd.DataFrame(counts).to_csv(ROOT / 'metadata/schmidt_count_qc.csv', index=False)
    return {'metadata_rows': len(meta), 'count_files': len(paths), 'genes_per_file': len(reference),
            'cell_lines': sorted(meta.Cell_line.unique().tolist()),
            'batch_labels': sorted(meta.Batch.unique().tolist()),
            'line_batch_combinations': len(meta[['Cell_line','Batch']].drop_duplicates()),
            'timepoint_d_missing_n': int(meta.timepoint_d.isna().sum()),
            'interpretation': 'Library counts are not independent biological replicate counts.'}


def read_block(sheet, header_row: int, start_col: int, end_col: int) -> pd.DataFrame:
    """Read an explicitly specified rectangle, excluding entirely empty rows."""
    headers = [sheet.cell(header_row, c).value for c in range(start_col, end_col + 1)]
    if any(h is None for h in headers) or len(set(headers)) != len(headers):
        raise ValueError(f'Invalid headers in {sheet.title}')
    rows = [[sheet.cell(r, c).value for c in range(start_col, end_col + 1)]
            for r in range(header_row + 1, sheet.max_row + 1)]
    return pd.DataFrame(rows, columns=headers).dropna(how='all')


def audit_lyons() -> dict:
    path = ROOT / 'data/raw/lyons/41467_2026_71962_MOESM4_ESM.xlsx'
    workbook = load_workbook(path, data_only=True)
    inventory, labels = [], []
    for sheet in workbook:
        inventory.append({'sheet': sheet.title, 'rows': sheet.max_row,
                          'columns': sheet.max_column, 'state': sheet.sheet_state,
                          'merged_ranges': ';'.join(str(r) for r in sheet.merged_cells.ranges)})
        # First five rows contain headers in most sheets; no numeric participant
        # measurements are copied to this dictionary.
        for row in sheet.iter_rows(max_row=min(5, sheet.max_row)):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip():
                    labels.append({'sheet': sheet.title, 'cell': cell.coordinate,
                                   'original_text': cell.value,
                                   'status': 'header candidate; not automatically a variable'})
    pd.DataFrame(inventory).to_csv(ROOT / 'metadata/lyons_sheet_inventory.csv', index=False)
    pd.DataFrame(labels).to_csv(ROOT / 'metadata/lyons_data_dictionary.csv', index=False)
    blocks = {}
    for name, start, end, keys in [('Figure 1a', 1, 4, ['Subject_ID','Dosing_day','Timepoint']),
                                  ('Figure 1b', 1, 4, ['Subject_ID','Dosing_day','Timepoint']),
                                  ('Table S4', 3, 7, ['ID','Dosing_day','Timepoint','Band']),
                                  ('Table S4', 10, 13, ['ID','Dosing_day','Timepoint'])]:
        frame = read_block(workbook[name], 3, start, end)
        require_unique(frame, keys)
        blocks[f'{name}:{start}-{end}'] = {
            'rows': len(frame), 'participants': int(frame[keys[0]].nunique()),
            'missing_by_column': {k: int(v) for k,v in frame.isna().sum().items()},
            'keys': keys,
            'dosing_labels': sorted(frame.Dosing_day.dropna().unique().tolist()),
            'timepoint_labels': sorted(frame.Timepoint.dropna().unique().tolist())}
    return {'sheets': len(workbook.sheetnames), 'identified_eeg_blocks': blocks,
            'warning': 'EEG identities do not establish outcome/intensity identities.'}


def main() -> None:
    (ROOT / 'metadata').mkdir(exist_ok=True)
    article_text = (ROOT / 'data/raw/lyons/article.html').read_text()
    article_status = ('Client Challenge; not article full text'
                      if 'Client Challenge' in article_text
                      else 'Response stored; semantic validation required')
    full_text = ROOT / 'data/raw/lyons/article_epmc.xml'
    if full_text.exists():
        document = ElementTree.fromstring(full_text.read_bytes())
        if not any(node.text == '10.1038/s41467-026-71962-3' for node in document.iter('article-id')):
            raise ValueError('Lyons full-text DOI mismatch')
        article_status = 'DOI-verified Europe PMC full-text XML available; original publisher challenge retained'
    summary = {'checksum_verified_downloads': verify_manifest(),
               'scope': 'Initial structural audit; not a completed feasibility or scientific analysis',
               'lyons_local_article_status': article_status,
               'schmidt': audit_schmidt(), 'lyons': audit_lyons()}
    (ROOT / 'metadata/audit_summary.json').write_text(json.dumps(summary, indent=2) + '\n')
    print(json.dumps(summary, indent=2))


if __name__ == '__main__':
    main()
